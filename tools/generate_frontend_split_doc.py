from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"F:\project\GuangFu\docs\光伏预测平台前端页面拆分文档.docx"


def set_run_font(run, size=None, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para_font(paragraph, size=11, bold=False, color=None):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_heading(doc, text, level):
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    set_run_font(r, size={1: 16, 2: 13, 3: 12}.get(level, 11), bold=True, color="2E74B5" if level < 3 else "1F4D78")
    p.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 10}.get(level, 6))
    p.paragraph_format.space_after = Pt({1: 10, 2: 7, 3: 5}.get(level, 4))
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, 11, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, 11)
    else:
        r = p.add_run(text)
        set_run_font(r, 11)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(item)
        set_run_font(r, 11)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], "E8EEF5")
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, 10.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i != 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r, 10)
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


def configure_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)


def main():
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    r = title.add_run("光伏预测平台前端页面拆分文档")
    set_run_font(r, 22, bold=True, color="0B2545")
    subtitle = doc.add_paragraph()
    sr = subtitle.add_run("依据《简要说明》整理，面向前端页面、模块、数据接口与交互落地。")
    set_run_font(sr, 11, color="555555")

    add_heading(doc, "1. 项目概览", 1)
    add_body(doc, "平台定位：面向光伏电站的预测、模型售卖、API 管理、报告分析与用户服务平台。核心使用场景是选定电站后查看预测结果、天气和服务器状态，并围绕模型与 API 完成试用、购买、创建、调用和统计。", "平台定位：")
    add_bullets(doc, [
        "数据来源包括长沙某电站 2023 年 6 月至今的光伏云图与功率数据、斯坦福大学 2017-2019 年光伏云图与功率数据。",
        "外部电站可从 pvoutput.org 选择，天气信息可接入当地气象局或天气服务。",
        "模型来源参考 OpenSTL、Time-Series-Library、Stanford solar forecasting dataset 以及相关多模态 irradiance forecasting 论文复现。",
    ])

    add_heading(doc, "2. 总体信息架构", 1)
    add_table(doc, ["端", "页面数量", "主要页面", "定位"], [
        ["PC 用户端", "8", "综合看板、模型选择、云图预测、模型广场、API 管理、综合分析报告、新闻通知、个人中心", "主业务入口，覆盖预测、模型、API、报告和账户"],
        ["管理端", "3", "用户/API 管理、电站管理、模型管理", "运营配置与后台管理，需要与用户端视觉明显区分"],
        ["小程序", "4", "API 状态、模型广场、个人信息、新闻", "只做用户端，复用 PC 核心能力并做移动化简化"],
    ], [1.1, 1.0, 2.5, 1.9])

    add_heading(doc, "3. PC 用户端页面拆分", 1)
    pc_rows = [
        ["U-01", "综合信息看板", "电站切换、预测表、天气、服务器状态", "选定电站、预测数据、天气、CPU/GPU/内存/硬盘", "电站选择器、预测表、天气卡片、服务器监控卡片", "P0"],
        ["U-02", "模型选择与使用", "按模型类别浏览并发起预测/试用", "模型列表、类别、模型说明、输入参数", "分类筛选、模型卡片、参数表单、运行结果区", "P0"],
        ["U-03", "云图预测", "上传 10 张云图并输出未来 10 张云图", "上传图片、SimVP+GSTA 预测结果", "图片上传墙、输出网格、进度状态、下载按钮", "P1"],
        ["U-04", "模型广场", "API 售卖与模型展示，支持购买/试用", "模型商品、价格、标签、调用说明", "模型卡片、详情抽屉、购买/试用、API 调用示例", "P0"],
        ["U-05", "API 管理", "API 创建、状态监测、使用统计", "API Key、调用量、余额、错误率、时延", "API 列表、创建弹窗、统计图表、调用日志", "P0"],
        ["U-06", "综合分析报告", "AI 对话式分析与报告导出", "电站数据、预测结果、用户问题、报告文件", "对话区、上下文选择、报告预览、导出按钮", "P0"],
        ["U-07", "新闻通知", "天气/行业新闻与通知展示", "中国天气网或新闻源、系统通知", "新闻列表、通知中心、已读状态", "P2"],
        ["U-08", "个人中心", "钱包、用户资料、安全设置", "用户资料、钱包余额、邮箱、密码", "资料表单、钱包记录、安全设置、绑定信息", "P1"],
    ]
    add_table(doc, ["编号", "页面", "页面目标", "数据依赖", "核心模块", "优先级"], pc_rows, [0.55, 1.05, 1.35, 1.35, 1.65, 0.55])

    add_heading(doc, "4. 关键页面功能说明", 1)
    page_specs = [
        ("U-01 综合信息看板", [
            "默认进入平台后的首页，顶部或侧栏提供电站切换。",
            "主区域展示预测表，建议按时间、预测功率/辐照度、置信状态、更新时间组织。",
            "右侧展示天气情况，包括天气、温度、湿度、风速、云量等字段。",
            "底部或侧栏展示服务器状态，包括 CPU、GPU、内存、硬盘，可用状态色区分健康/告警。",
        ]),
        ("U-02 模型选择与使用", [
            "一级分类：时序模型、视频时空递归模型、多模态模型。",
            "二级内容展示分类下具体模型，并提供模型说明、输入要求、输出说明和运行入口。",
            "时序模型预测逻辑暂按读取过去 30 分钟、预测未来 6 个步长、每步 5 分钟设计，后续可通过配置调整。",
        ]),
        ("U-03 云图预测", [
            "只接入一个云图预测模型：SimVP+GSTA。",
            "上传区限制 10 张云图，输出区展示 10 张预测云图。",
            "由于实时云图获取受限，该页面更适合作为离线上传预测或模型能力演示。",
        ]),
        ("U-04 模型广场", [
            "参考硅基流动模型广场的视觉和交互，强调卡片展示、标签、价格/额度和调用说明。",
            "不能只做静态列表，需要实现购买/试用/查看 API 调用方式等功能闭环。",
            "模型详情页或抽屉展示模型介绍、输入输出、计费方式、调用示例和相关指标。",
        ]),
        ("U-05 API 管理", [
            "将 API 创建、Key 管理、调用统计、使用量趋势、调用日志集中在一个页面中。",
            "列表建议一条 API 一行，字段包括 API 名称、所属模型、Key 状态、今日调用量、总调用量、错误率、创建时间、操作。",
            "提供创建、启停、复制 Key、重置 Key、查看统计、删除等操作。",
        ]),
        ("U-06 综合分析报告", [
            "采用 AI 网页对话格式，左侧可放历史报告/会话，主区域为对话与报告生成。",
            "可接入自有 API、Agent 或微调小模型，用于分析电站预测结果和运行情况。",
            "必须支持分析报告导出，建议支持 docx/pdf 两种导出目标。",
        ]),
    ]
    for name, items in page_specs:
        add_heading(doc, name, 2)
        add_bullets(doc, items)

    add_heading(doc, "5. 模型分类与页面展示", 1)
    add_table(doc, ["模型类别", "模型", "前端展示建议"], [
        ["时序基线", "PatchTST、DLinear、iTransformer、TimeXer、TimeMixer、TSMixer、Transformer", "作为时序预测模型集合，适合卡片组+参数配置+运行结果曲线。"],
        ["云图/视觉融合", "CNN+MLP、CNN+LSTM、3D-CNN+LSTM、convlstm+lstm", "突出云图输入能力，展示输入样例、视觉特征和预测结果。"],
        ["视频时空递归", "simvp+gsta、TAU、PredRNN、PredRNN++、ConvLSTM、E3D-LSTM、swinLSTM、sunset", "适合展示连续帧输入/输出，支持帧序列预览。"],
    ], [1.25, 3.05, 2.2])

    add_heading(doc, "6. 管理端页面拆分", 1)
    add_table(doc, ["编号", "页面", "管理对象", "核心功能", "备注"], [
        ["A-01", "用户/API 管理", "用户账号与用户 API", "用户查询、API 状态、调用量、余额/额度、禁用/恢复、日志查看", "以 API 信息为中心，一条 API 对应一行"],
        ["A-02", "电站管理", "用户端首页所针对的电站", "新增电站、编辑位置、数据源配置、天气源配置、启停状态", "与首页电站切换联动"],
        ["A-03", "模型管理", "模型广场中的模型", "模型名称、介绍、类别、上下架、增添、删除、价格/额度配置", "影响模型广场展示"],
    ], [0.55, 1.1, 1.35, 2.55, 0.95])
    add_body(doc, "管理端视觉建议：与用户端明显区分，使用更偏后台系统的布局，如左侧导航、顶部搜索/操作区、数据表格为主，减少营销化卡片。", "管理端视觉建议：")

    add_heading(doc, "7. 小程序页面拆分", 1)
    add_table(doc, ["编号", "页面", "复用来源", "移动端调整"], [
        ["M-01", "用户 API 状态监测", "PC U-05 API 管理", "只保留 API 状态、调用量、余额、告警，不放复杂创建流程"],
        ["M-02", "模型广场", "PC U-04 模型广场", "卡片单列展示，详情页承载购买/试用与调用说明"],
        ["M-03", "用户个人信息", "PC U-08 个人中心", "资料、钱包、安全设置分组折叠"],
        ["M-04", "新闻通知", "PC U-07 新闻通知", "列表+详情，支持未读提示"],
    ], [0.55, 1.2, 1.55, 3.2])

    add_heading(doc, "8. 公共组件与接口对象", 1)
    add_table(doc, ["组件/对象", "复用页面", "说明"], [
        ["电站选择器", "U-01、A-02", "选择开源平台或自建电站，驱动预测表、天气和报告上下文。"],
        ["模型卡片", "U-02、U-04、A-03、M-02", "展示模型名称、类别、简介、状态、价格、操作入口。"],
        ["API Key 表格", "U-05、A-01、M-01", "围绕 API 一行一条记录，支持状态和统计查看。"],
        ["预测结果图表", "U-01、U-02、U-06", "展示时间序列预测、实际值对比、置信状态。"],
        ["图片上传/预览", "U-03", "限制 10 张输入图，输出 10 张预测图。"],
        ["报告导出", "U-06", "导出综合分析报告，建议预留 PDF/DOCX。"],
    ], [1.35, 1.6, 3.55])

    add_heading(doc, "9. 前端路由建议", 1)
    add_table(doc, ["端", "路由", "页面"], [
        ["PC", "/dashboard", "综合信息看板"],
        ["PC", "/models/use", "模型选择与使用"],
        ["PC", "/cloud-forecast", "云图预测"],
        ["PC", "/marketplace", "模型广场"],
        ["PC", "/api", "API 管理"],
        ["PC", "/reports", "综合分析报告"],
        ["PC", "/news", "新闻通知"],
        ["PC", "/profile", "个人中心"],
        ["Admin", "/admin/users-apis", "用户/API 管理"],
        ["Admin", "/admin/stations", "电站管理"],
        ["Admin", "/admin/models", "模型管理"],
    ], [0.9, 2.25, 3.35])

    add_heading(doc, "10. 实施优先级", 1)
    add_table(doc, ["优先级", "页面/模块", "理由"], [
        ["P0", "综合信息看板、模型选择与使用、模型广场、API 管理、综合分析报告", "构成核心业务闭环：预测展示、模型使用、模型售卖、API 统计、报告生成。"],
        ["P1", "云图预测、个人中心、管理端三页", "支撑演示亮点、账号体系和运营维护。"],
        ["P2", "新闻通知、小程序完整适配", "偏装饰或移动端增强，可在核心链路稳定后推进。"],
    ], [0.8, 2.2, 3.5])

    add_heading(doc, "11. 待确认事项", 1)
    add_bullets(doc, [
        "时序预测输入窗口和预测步长是否最终固定，或需要在模型配置中动态维护。",
        "模型广场是否需要真实支付/钱包扣费，还是仅实现额度购买和试用流程。",
        "API 管理中是否需要区分用户自建 API、购买模型 API、系统内置 API。",
        "综合分析报告导出的格式、模板和报告字段需要进一步定义。",
        "天气、新闻、电站开源平台的数据接口稳定性和可用权限需要提前验证。",
    ])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("光伏预测平台前端拆分")
    set_run_font(fr, 9, color="555555")

    doc.save(OUT)


if __name__ == "__main__":
    main()
